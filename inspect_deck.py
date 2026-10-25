from pptx import Presentation
prs = Presentation(r"C:\Users\wpbro\AppData\Local\Temp\deck_work.pptx")
print("Slides:", len(prs.slides))
print("Layouts:", len(prs.slide_layouts))
for i, sl in enumerate(prs.slides):
    shapes_txt = [s for s in sl.shapes if s.has_text_frame]
    first = shapes_txt[0].text_frame.text[:70] if shapes_txt else "(no text)"
    print(f"  Slide {i+1}: {len(sl.shapes)} shapes | {repr(first)}")
